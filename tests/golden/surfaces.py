"""Every user-facing surface M10a touches, rendered from committed inputs.

The golden *cycle* protects the pipeline. This protects the other half: Step 6
restructures ``StatsReport`` and Step 7 edits six commands, and "crypto's output is
unchanged" is a claim about what the owner reads, not only about what the gate
decides. So each surface is rendered here, before any edit, and pinned.

Three rules this module keeps:

* **Real read models, real ORM rows.** ``stats_view``, ``pulse_view``,
  ``position_view``, ``snapshot_view``, ``build_journal`` and the card functions are
  the production ones, driven by real :mod:`sentinel.storage.models` rows — the same
  discipline ``tests/bot/test_journal.py`` and ``test_pulse.py`` already keep, so a
  renamed column fails here rather than on somebody's phone.
* **Populated, not empty.** An empty view renders an "and nothing to show" branch
  and would pin almost nothing. Every surface below is given rows on every branch it
  has: both spend states, a filled and an unfilled position, all four journal
  populations, a capped field, a degraded snapshot.
* **Frozen everything.** One instant, one timezone, fixed ids. Nothing here reads a
  clock or a database.
"""

from __future__ import annotations

import io
import json
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from typing import Any
from uuid import UUID

from openpyxl import load_workbook

from sentinel.bot.cards import (
    positions_card,
    pulse_card,
    pulse_day_card,
    snapshot_card,
    stats_card,
    status_card,
    symbol_pulse_card,
    watchlist_card,
)
from sentinel.bot.export import journal_workbook
from sentinel.bot.markets import section_header
from sentinel.bot.pulse import pulse_day_view, pulse_view, symbol_pulse_view
from sentinel.bot.readmodels import position_view, spend_view, stats_view
from sentinel.bot.snapshot import snapshot_view
from sentinel.bot.views import DataSourceView, StatusView
from sentinel.core.config import AppConfig, load_config
from sentinel.core.markets import Market
from sentinel.core.pauses import effective_pause
from sentinel.llm.spend import SpendTotals
from sentinel.risk.models import PauseReason, PauseState
from sentinel.screener.models import DirectionHint, ScreenerVerdict
from sentinel.stats.compute import by_key, split, summarize, tracked
from sentinel.stats.journal import build_journal
from sentinel.stats.models import Book, BookStats, Population, StatsReport
from sentinel.stats.queries import resolved_from_row
from sentinel.storage.models import (
    AnalystReportRow,
    CycleRow,
    GateDecisionRow,
    MarketSnapshotRow,
    SignalExitRow,
    SignalFillRow,
    SignalRow,
)
from tests.golden.pipeline import TZ, golden_features, golden_snapshot
from tests.risk_double import approved_plan

#: One instant for every surface. August, so Europe/Vilnius is UTC+3 — a real
#: offset rather than zero, which is what makes a UTC-only bug visible.
NOW = datetime(2026, 8, 18, 12, 0, tzinfo=UTC)

OWNER_ID = 7222549221
CYCLE_ID = UUID("11111111-1111-4111-8111-111111111111")

WINDOW = "30d"
SINCE = NOW - timedelta(days=30)


def _uuid(tag: int) -> UUID:
    return UUID(f"{tag:08d}-0000-4000-8000-000000000000")


# --------------------------------------------------------------------------- #
# Signals — the rows /stats, /positions and /journal all read
# --------------------------------------------------------------------------- #


def _signal(
    tag: int,
    *,
    decision: str | None,
    dry_run: bool = False,
    realized_r: str | None = None,
    outcome: str | None = None,
    filled: bool = True,
    tp_hits: int = 0,
    setup_type: str = "trend_pullback",
    plan: dict[str, Any] | None = None,
) -> SignalRow:
    """One ``signals`` row, detached from any session."""
    row = SignalRow(
        id=_uuid(tag),
        plan_id=_uuid(1000 + tag),
        cycle_id=CYCLE_ID,
        user_id=OWNER_ID,
        # Set explicitly: a server default fills the column in Postgres, not on a
        # detached instance, and a row with market=None is not a row the database
        # can hold. Every fixture here is what a stored row really looks like.
        market=Market.CRYPTO.value,
        symbol="SOLUSDT",
        direction="long",
        setup_type=setup_type,
        prompt_version="fable_v1",
        confidence=78,
        created_at=NOW - timedelta(days=tag),
        expires_at=NOW - timedelta(days=tag) + timedelta(hours=12),
        status="CLOSED" if outcome else "PENDING_ENTRY",
        decision=decision,
        decided_at=None if decision is None else NOW - timedelta(days=tag),
        decided_by_user_id=None if decision is None else OWNER_ID,
        plan=plan or {},
        chart_params=[],
        dry_run=dry_run,
        filled_qty=Decimal("1.5") if filled else Decimal("0"),
        avg_fill_price=Decimal("82.55") if filled else None,
        tp_hits=tp_hits,
        outcome=outcome,
        closed_at=None if outcome is None else NOW - timedelta(days=tag, hours=-6),
    )
    # `number` is an IDENTITY column: Postgres assigns it, so a detached row has
    # none. Every surface prints it, so it is set explicitly rather than left None.
    row.number = tag
    if realized_r is not None:
        row.realized_r = Decimal(realized_r)
        row.realized_eur = Decimal(realized_r) * Decimal("75")
        row.realized_costs_eur = Decimal("3.16")
    return row


def signal_rows(config: AppConfig) -> list[SignalRow]:
    """A book with every population and every outcome branch populated.

    Deliberately includes a win, a loss, a scratch (a stop moved to breakeven, which
    is neither), an unfilled expiry, a dry-run row and an undecided one — those are
    the six branches ``compute.summarize`` and the journal's four sheets have between
    them, and a golden that exercised three of them would pin very little.
    """
    plan = approved_plan(config).model_dump(mode="json")
    return [
        _signal(1, decision="TAKEN", realized_r="1.60", outcome="TP2", tp_hits=2, plan=plan),
        _signal(2, decision="TAKEN", realized_r="-1.00", outcome="STOP", plan=plan),
        _signal(3, decision="TAKEN", realized_r="0.00", outcome="STOP", tp_hits=1, plan=plan),
        _signal(
            4,
            decision="WATCHING",
            realized_r="0.78",
            outcome="TP1",
            tp_hits=1,
            setup_type="breakout_retest",
            plan=plan,
        ),
        _signal(5, decision="SKIPPED", realized_r="-1.00", outcome="STOP", plan=plan),
        _signal(6, decision="WATCHING", outcome="EXPIRY", filled=False, plan=plan),
        _signal(7, decision=None, plan=plan),
        _signal(
            8,
            decision="TAKEN",
            dry_run=True,
            realized_r="0.40",
            outcome="TP1",
            tp_hits=1,
            plan=plan,
        ),
    ]


# --------------------------------------------------------------------------- #
# /stats
# --------------------------------------------------------------------------- #


def stats_report(config: AppConfig) -> StatsReport:
    """What ``build_report`` assembles, without needing a session to assemble it."""
    resolved = [resolved_from_row(row) for row in signal_rows(config) if row.outcome]
    books = split(resolved)
    followed = tracked(resolved)
    return StatsReport(
        window=WINDOW,
        since=SINCE,
        generated_at=NOW,
        market=Market.CRYPTO,
        books=tuple(
            BookStats(
                book=Book(population=population, market=Market.CRYPTO),
                stats=summarize(books[Book(population=population, market=Market.CRYPTO)]),
            )
            for population in Population
        ),
        by_setup=by_key(followed, "setup_type"),
        by_prompt_version=by_key(followed, "prompt_version"),
    )


# --------------------------------------------------------------------------- #
# /pulse
# --------------------------------------------------------------------------- #


def _cycle(*, cycle_id: UUID = CYCLE_ID, dry_run: bool = False) -> CycleRow:
    row = CycleRow(
        cycle_id=cycle_id,
        market=Market.CRYPTO.value,
        started_at=NOW - timedelta(minutes=6),
        finished_at=NOW - timedelta(minutes=2),
        status="OK",
        dry_run=dry_run,
        symbols_requested=10,
        symbols_scanned=10,
        symbols_skipped=2,
        skipped={
            "LINKUSDT": {
                "reason": "RECENTLY_ANALYSED",
                "detail": "analysed since 2026-08-18T11:00:00+00:00 and not a candidate then",
            },
            "ADAUSDT": {
                "reason": "OPEN_SIGNAL",
                "detail": "a signal is already open for this symbol (PRD F11)",
            },
        },
        candidates=3,
        analyzed=1,
        approved=1,
        published=1,
        spend_usd_estimate=Decimal("0.31428000"),
        analysis_suspended=False,
        suspended_reason=None,
        error=None,
    )
    return row


def _screener() -> tuple[ScreenerVerdict, ...]:
    return (
        ScreenerVerdict(
            symbol="SOLUSDT",
            interesting=True,
            direction_hint=DirectionHint.LONG,
            reason="1h pullback into EMA50 with volume drying up; 4h trend intact.",
        ),
        ScreenerVerdict(
            symbol="LINKUSDT",
            interesting=True,
            direction_hint=DirectionHint.UNCLEAR,
            reason="Range compression at the top of the 4h box.",
        ),
        ScreenerVerdict(
            symbol="BTCUSDT",
            interesting=False,
            direction_hint=DirectionHint.UNCLEAR,
            reason="Mid-range, no level in play.",
        ),
    )


def _report_row(config: AppConfig) -> AnalystReportRow:
    """A stored analyst verdict, with the JSONB the drill-down card reads back."""
    report = approved_plan(config).report
    payload = report.model_dump(mode="json")
    row = AnalystReportRow(
        id=_uuid(9001),
        cycle_id=CYCLE_ID,
        snapshot_id=_uuid(9002),
        llm_call_id=_uuid(9003),
        market=Market.CRYPTO.value,
        symbol="SOLUSDT",
        created_at=NOW - timedelta(minutes=4),
        role="primary",
        provider="fable5",
        model="claude-fable-5",
        prompt_version="fable_v1",
        candidate_status=report.candidate_status.value,
        setup_type=report.setup_type.value,
        direction=report.direction.value,
        confidence=report.confidence,
        thesis=report.thesis,
        report=json.loads(json.dumps(payload)),
    )
    return row


def _gate_rows() -> tuple[GateDecisionRow, ...]:
    """One approval and one rejection, so both gate wordings are pinned.

    The rejection is ``NET_RR_TOO_LOW`` — a PERSONAL code under M8.4's shared/personal
    split, so the pulse card's privacy branch is exercised rather than only its happy
    path.
    """
    return (
        GateDecisionRow(
            id=_uuid(9101),
            cycle_id=CYCLE_ID,
            user_id=OWNER_ID,
            market=Market.CRYPTO.value,
            symbol="SOLUSDT",
            evaluated_at=NOW - timedelta(minutes=3),
            gate_status="APPROVED_FOR_HUMAN",
            reason=None,
            message="approved for human review",
            prompt_version="fable_v1",
            plan=None,
        ),
        GateDecisionRow(
            id=_uuid(9102),
            cycle_id=CYCLE_ID,
            user_id=OWNER_ID,
            market=Market.CRYPTO.value,
            symbol="LINKUSDT",
            evaluated_at=NOW - timedelta(minutes=3),
            gate_status="REJECTED",
            reason="NET_RR_TOO_LOW",
            message="reward-to-risk at TP1 is below the minimum once fees and funding are paid",
            prompt_version="fable_v1",
            plan=None,
        ),
    )


def _spend(config: AppConfig) -> Any:
    """The owner's spend line, in the WARN state so its wording is pinned too.

    **Takes the config since M10d; it called ``load_config()`` directly before.** That
    made ``test_the_deployed_config_renders_the_same_surface`` blind to the whole
    ``llm:`` block: both sides of that comparison rendered the *shipped* limits
    whichever config they were handed, so the two shapes agreed by construction on
    every figure this line prints. Found when raising the ceiling moved ``/pulse`` and
    the legacy comparison did not notice. A helper that ignores its own parameter is
    the same shape of defect as a guard that compares strings where it means values
    (journal/M10c_REPORT.md §3): it reads like coverage.
    """
    return spend_view(
        SpendTotals(
            day_usd=Decimal("7.42"),
            month_usd=Decimal("38.10"),
            calls=214,
            unpriced_calls=0,
        ),
        config.llm,
    )


# --------------------------------------------------------------------------- #
# /positions
# --------------------------------------------------------------------------- #


def _position_rows(config: AppConfig) -> tuple[Any, ...]:
    plan = approved_plan(config)
    row = _signal(1, decision="TAKEN", plan=plan.model_dump(mode="json"))
    row.status = "PARTIALLY_FILLED"
    row.stop_price_current = plan.entries[0].price
    row.tp_hits = 1

    fills = (
        SignalFillRow(
            id=1,
            signal_id=row.id,
            rung_index=0,
            price=plan.entries[0].price,
            qty=plan.entries[0].qty,
            filled_at=NOW - timedelta(hours=3),
            detected_at=NOW - timedelta(hours=3),
            source="tracker",
        ),
    )
    exits = (
        SignalExitRow(
            id=1,
            signal_id=row.id,
            kind="TP1",
            price=plan.targets[0],
            qty=plan.entries[0].qty / Decimal("2"),
            exited_at=NOW - timedelta(hours=1),
            detected_at=NOW - timedelta(hours=1),
        ),
    )

    unfilled = _signal(2, decision="TAKEN", filled=False, plan=plan.model_dump(mode="json"))
    unfilled.status = "PENDING_ENTRY"
    unfilled.filled_qty = Decimal("0")

    return (
        position_view(row, plan, fills, exits, mark_price=plan.targets[0]),
        position_view(unfilled, plan, (), (), mark_price=None),
    )


# --------------------------------------------------------------------------- #
# /snapshot
# --------------------------------------------------------------------------- #


def _snapshot_row(config: AppConfig) -> MarketSnapshotRow:
    """A stored snapshot carrying the real computed features from the cassettes.

    Marked DEGRADED with two fields named, because the degraded branch is the one
    that renders the warning line — and because the EMA-stack label that goes through
    it is the value that made ``/snapshot`` silent in M8.6.
    """
    snapshot = golden_snapshot()
    features = golden_features(snapshot, config)
    return MarketSnapshotRow(
        id=_uuid(9201),
        cycle_id=CYCLE_ID,
        market=Market.CRYPTO.value,
        symbol=snapshot.symbol,
        captured_at=snapshot.captured_at,
        schema_version=1,
        last_price=snapshot.last_price,
        data_quality="DEGRADED",
        degraded_fields=["funding", "news"],
        context={"features": json.loads(json.dumps(features.model_dump(mode="json")))},
        sources={},
    )


# --------------------------------------------------------------------------- #
# /journal
# --------------------------------------------------------------------------- #


def _journal_json(config: AppConfig) -> dict[str, Any]:
    """The workbook read back through openpyxl: sheets, headers, and one row each.

    Read back rather than asserted on the builder's output, because the file is what
    leaves the process. Sheet **order** is captured as a list, so a reordering fails
    even when every sheet still exists.
    """
    rows = signal_rows(config)
    books = build_journal(rows, fills={}, exits={})
    payload = journal_workbook(books, window_label="last 30 days", generated_at=NOW, tz=TZ)

    workbook = load_workbook(io.BytesIO(payload))
    sheets: dict[str, Any] = {}
    for name in workbook.sheetnames:
        sheet = workbook[name]
        values = list(sheet.iter_rows(values_only=True))
        sheets[name] = {
            "header": [
                str(cell) if cell is not None else None for cell in (values[0] if values else ())
            ],
            "first_row": [
                cell.isoformat() if isinstance(cell, datetime) else cell
                for cell in (values[1] if len(values) > 1 else ())
            ],
            "row_count": len(values),
        }
    return {"order": list(workbook.sheetnames), "active": workbook.active.title, "sheets": sheets}


def _status_view(config: AppConfig) -> StatusView:
    """``/status`` with every block populated, and a pause in force.

    Included because M10a rewrites how the pause is *composed* — four sources
    instead of two, through ``core/pauses.py`` — and the scope label is the one part
    of this card the change can move. A paused card is the one worth pinning; the
    unpaused branch is one line.
    """
    pause = effective_pause(
        global_pause=PauseState(paused=True, reason=PauseReason.MANUAL, until=None),
        market_pause=PauseState(),
        user_pause=PauseState(),
        user_market_pause=PauseState(),
        market=Market.CRYPTO,
        now=NOW,
    )
    return StatusView(
        paused=pause.active,
        pause_reason=None if pause.state.reason is None else pause.state.reason.value,
        paused_until=pause.state.until,
        pause_scope=pause.label(multi_market=config.multi_market),
        market=section_header(Market.CRYPTO, config),
        capital_eur=Decimal("10000"),
        risk_per_trade_pct=config.risk.risk_per_trade_pct,
        watchlist_size=len(config.market(Market.CRYPTO).watchlist),
        signals_total=8,
        signals_undecided=1,
        signals_taken=4,
        stuck_messages=1,
        data_sources=(
            DataSourceView(
                symbol="BTCUSDT",
                quality="OK",
                captured_at=NOW - timedelta(minutes=4),
            ),
            DataSourceView(
                symbol="SOLUSDT",
                quality="DEGRADED",
                captured_at=NOW - timedelta(minutes=4),
                degraded_fields=("funding", "news"),
            ),
        ),
        dry_run=config.market(Market.CRYPTO).dry_run,
        last_cycle_at=NOW - timedelta(minutes=2),
        last_cycle_status="OK",
        cycles_completed=23,
        cycles_started=24,
        open_risk_pct=Decimal("1.50"),
        max_open_risk_pct=config.risk.max_open_risk_pct,
        open_positions=2,
        max_positions=config.risk.max_positions,
        signals_today=2,
        max_signals_per_day=config.risk.max_signals_per_day,
        signals_open=2,
        spend=spend_view(
            SpendTotals(day_usd=Decimal("7.42"), month_usd=Decimal("38.10"), calls=214),
            config.llm,
            market=config.market(Market.CRYPTO),
        ),
    )


# --------------------------------------------------------------------------- #
# The whole set
# --------------------------------------------------------------------------- #

#: The page separator used when a card renders as several messages. A literal that
#: cannot occur in Telegram HTML, so a page-count change fails the golden instead of
#: being absorbed into the joined text.
PAGE_BREAK = "\n\n=== PAGE BREAK ===\n\n"


def render_surfaces(config: AppConfig | None = None) -> dict[str, Any]:
    """Every pinned surface, keyed by golden filename.

    ``config`` defaults to the repo's ``config.yaml``. It is a parameter so the same
    render can be driven from the **frozen pre-M10a config the server actually runs**
    — see ``test_golden_surfaces.py``, where the two are asserted to be identical.
    """
    config = config or load_config()

    cycle = _cycle()
    screener = _screener()
    report_row = _report_row(config)
    gate_rows = _gate_rows()
    spend = _spend(config)

    pulse = pulse_view(
        cycle,
        screener=screener,
        reports=[report_row],
        decisions=list(gate_rows),
        spend=spend,
    )
    day = pulse_day_view(
        [cycle, _cycle(cycle_id=_uuid(9301), dry_run=True)],
        since=NOW - timedelta(hours=24),
        started=24,
        screener={CYCLE_ID: screener, _uuid(9301): screener},
        reports=[report_row],
        decisions=list(gate_rows),
        spend=spend,
    )
    symbol = symbol_pulse_view(report_row, list(gate_rows), symbol="SOLUSDT", on_watchlist=True)

    return {
        "status": status_card(_status_view(config), TZ),
        "stats": stats_card(
            stats_view(stats_report(config), header=section_header(Market.CRYPTO, config)),
            TZ,
        ),
        # ``header=`` matters, and was missing until M10c. The real handlers pass
        # ``section_header`` (bot/handlers/commands.py:430 and :459) and this rendered
        # without it — so with one market the golden was right by coincidence, and the
        # multi-market golden below would have under-pinned exactly the line switching
        # forex on adds. With one market ``section_header`` returns "", so adding it
        # moves no existing byte; the sha256 of pulse.txt and pulse_24h.txt are
        # unchanged by this commit, which is the check that says so.
        "pulse": pulse_card(pulse, TZ, header=section_header(Market.CRYPTO, config)),
        "pulse_24h": pulse_day_card(day, TZ, header=section_header(Market.CRYPTO, config)),
        "pulse_symbol": PAGE_BREAK.join(symbol_pulse_card(symbol, TZ)),
        "positions": positions_card(
            _position_rows(config), TZ, header=section_header(Market.CRYPTO, config)
        ),
        "watchlist": watchlist_card(
            config.market(Market.CRYPTO).watchlist,
            "yaml",
            header=section_header(Market.CRYPTO, config),
        ),
        "snapshot": snapshot_card(
            snapshot_view(_snapshot_row(config), symbol="BTCUSDT", on_watchlist=True), TZ
        ),
        "journal": _journal_json(config),
    }


__all__ = ["NOW", "OWNER_ID", "PAGE_BREAK", "render_surfaces", "signal_rows", "stats_report"]
